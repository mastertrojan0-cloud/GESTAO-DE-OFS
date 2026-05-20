"""
metrics.py — API de metricas do Sistema OFS.
Endpoints para dashboard, graficos e exportacao.
"""

import io
import re
from datetime import datetime, date, timezone, timedelta
from typing import Optional
from uuid import UUID

try:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    _HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    _HEADER_FILL = PatternFill(start_color="1A1A1A", end_color="1A1A1A", fill_type="solid")
    _HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
    _CELL_ALIGN = Alignment(horizontal="left", vertical="center")
    _THIN_BORDER = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
except ImportError:
    Font = None
    PatternFill = None
    Alignment = None
    Border = None
    Side = None
    get_column_letter = None
    _HEADER_FONT = None
    _HEADER_FILL = None
    _HEADER_ALIGN = None
    _CELL_ALIGN = None
    _THIN_BORDER = None

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from dateutil.relativedelta import relativedelta

from app.db.session import get_db
from app.core.security import get_current_user, require_role
from app.core.audit import audit
from app.models.usuario import Usuario
from app.models.empresa import Empresa
from app.models.meta import Meta
from app.schemas.metrics import (
    WeeklyMetricsResponse,
    ProgrammedVsRealizedResponse,
    PositiveVsNegativeResponse,
    ByCompanyResponse,
    ByUserResponse,
    ByShiftResponse,
    TopBehaviorsResponse,
    WeeklyEvolutionResponse,
)
from app.services.metrics_service import (
    calculate_weekly_metrics,
    get_weekly_evolution,
    get_company_ranking,
    get_user_ranking,
    get_shift_distribution,
    get_top_behaviors,
    get_programmed_vs_realized,
    get_positive_vs_negative,
    invalidate_metrics_cache,
)

router = APIRouter(prefix="/api/metrics", tags=["Metricas"])

# Constantes
ALL_ROLES = ["observador", "supervisor", "gestor", "admin"]
STAF_ROLES = ["supervisor", "gestor", "admin"]
DEFAULT_WEEKS = 4


def _parse_date(value: str) -> date:
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", value):
        raise HTTPException(status_code=400, detail=f"Data invalida: {value}. Use YYYY-MM-DD.")
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=400, detail=f"Data invalida: {value}. Use YYYY-MM-DD.")


def _default_week_range() -> tuple[date, date]:
    hoje = date.today()
    inicio = hoje - timedelta(days=hoje.weekday())
    segunda = inicio if inicio.weekday() == 0 else inicio + timedelta(days=7) - timedelta(days=hoje.weekday())
    segunda = hoje - timedelta(days=hoje.weekday())
    domingo = segunda + timedelta(days=6)
    return segunda, domingo


# ═══════════════════════════════════════════════════════════════════════
# GET /api/metrics/weekly
# ═══════════════════════════════════════════════════════════════════════

@router.get("/weekly", response_model=WeeklyMetricsResponse)
async def get_weekly_metrics(
    data_inicio: str = Query(default=None, description="YYYY-MM-DD"),
    data_fim: str = Query(default=None, description="YYYY-MM-DD"),
    empresa_id: Optional[UUID] = Query(None),
    current_user: Usuario = Depends(require_role(*ALL_ROLES)),
    db: AsyncSession = Depends(get_db),
):
    inicio = _parse_date(data_inicio) if data_inicio else _default_week_range()[0]
    fim = _parse_date(data_fim) if data_fim else _default_week_range()[1]

    if inicio > fim:
        raise HTTPException(status_code=400, detail="data_inicio deve ser anterior a data_fim.")

    try:
        year = inicio.isocalendar()[0]
        week = inicio.isocalendar()[1]
        metrics = await calculate_weekly_metrics(db, year, week, empresa_id, current_user)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erro ao calcular metricas: {str(e)}")

    if metrics["ofc_realizadas"] == 0:
        return WeeklyMetricsResponse(**{
            **metrics,
            "status": "SEM_DADOS",
            "indicadores": metrics.get("indicadores", []),
        })

    return WeeklyMetricsResponse(**metrics)


# ═══════════════════════════════════════════════════════════════════════
# GET /api/metrics/charts/programmed-vs-realized
# ═══════════════════════════════════════════════════════════════════════

@router.get("/charts/programmed-vs-realized", response_model=ProgrammedVsRealizedResponse)
async def chart_programmed_vs_realized(
    data_inicio: str = Query(..., description="YYYY-MM-DD"),
    data_fim: str = Query(..., description="YYYY-MM-DD"),
    empresa_id: Optional[UUID] = Query(None),
    current_user: Usuario = Depends(require_role(*ALL_ROLES)),
    db: AsyncSession = Depends(get_db),
):
    inicio = _parse_date(data_inicio)
    fim = _parse_date(data_fim)

    data = await get_programmed_vs_realized(db, inicio, fim, empresa_id, current_user)
    return ProgrammedVsRealizedResponse(chart=data)


# ═══════════════════════════════════════════════════════════════════════
# GET /api/metrics/charts/positive-vs-negative
# ═══════════════════════════════════════════════════════════════════════

@router.get("/charts/positive-vs-negative", response_model=PositiveVsNegativeResponse)
async def chart_positive_vs_negative(
    data_inicio: str = Query(..., description="YYYY-MM-DD"),
    data_fim: str = Query(..., description="YYYY-MM-DD"),
    empresa_id: Optional[UUID] = Query(None),
    current_user: Usuario = Depends(require_role(*ALL_ROLES)),
    db: AsyncSession = Depends(get_db),
):
    inicio = _parse_date(data_inicio)
    fim = _parse_date(data_fim)

    data = await get_positive_vs_negative(db, inicio, fim, empresa_id, current_user)
    return PositiveVsNegativeResponse(**data)


# ═══════════════════════════════════════════════════════════════════════
# GET /api/metrics/charts/by-company
# ═══════════════════════════════════════════════════════════════════════

@router.get("/charts/by-company", response_model=ByCompanyResponse)
async def chart_by_company(
    data_inicio: str = Query(..., description="YYYY-MM-DD"),
    data_fim: str = Query(..., description="YYYY-MM-DD"),
    empresa_id: Optional[UUID] = Query(None),
    current_user: Usuario = Depends(require_role(*ALL_ROLES)),
    db: AsyncSession = Depends(get_db),
):
    inicio = _parse_date(data_inicio)
    fim = _parse_date(data_fim)

    data = await get_company_ranking(db, inicio, fim, empresa_id, current_user)
    return ByCompanyResponse(chart=data)


# ═══════════════════════════════════════════════════════════════════════
# GET /api/metrics/charts/by-user
# ═══════════════════════════════════════════════════════════════════════

@router.get("/charts/by-user", response_model=ByUserResponse)
async def chart_by_user(
    data_inicio: str = Query(..., description="YYYY-MM-DD"),
    data_fim: str = Query(..., description="YYYY-MM-DD"),
    empresa_id: Optional[UUID] = Query(None),
    current_user: Usuario = Depends(require_role(*ALL_ROLES)),
    db: AsyncSession = Depends(get_db),
):
    inicio = _parse_date(data_inicio)
    fim = _parse_date(data_fim)

    data = await get_user_ranking(db, inicio, fim, empresa_id, current_user)
    return ByUserResponse(chart=data)


# ═══════════════════════════════════════════════════════════════════════
# GET /api/metrics/charts/by-shift
# ═══════════════════════════════════════════════════════════════════════

@router.get("/charts/by-shift", response_model=ByShiftResponse)
async def chart_by_shift(
    data_inicio: str = Query(..., description="YYYY-MM-DD"),
    data_fim: str = Query(..., description="YYYY-MM-DD"),
    empresa_id: Optional[UUID] = Query(None),
    current_user: Usuario = Depends(require_role(*ALL_ROLES)),
    db: AsyncSession = Depends(get_db),
):
    inicio = _parse_date(data_inicio)
    fim = _parse_date(data_fim)

    data = await get_shift_distribution(db, inicio, fim, empresa_id, current_user)
    return ByShiftResponse(chart=data)


# ═══════════════════════════════════════════════════════════════════════
# GET /api/metrics/charts/top-behaviors
# ═══════════════════════════════════════════════════════════════════════

@router.get("/charts/top-behaviors", response_model=TopBehaviorsResponse)
async def chart_top_behaviors(
    data_inicio: str = Query(..., description="YYYY-MM-DD"),
    data_fim: str = Query(..., description="YYYY-MM-DD"),
    empresa_id: Optional[UUID] = Query(None),
    limit: int = Query(10, ge=1, le=50),
    current_user: Usuario = Depends(require_role(*ALL_ROLES)),
    db: AsyncSession = Depends(get_db),
):
    inicio = _parse_date(data_inicio)
    fim = _parse_date(data_fim)

    data = await get_top_behaviors(db, inicio, fim, empresa_id, current_user, limit)
    return TopBehaviorsResponse(**data)


# ═══════════════════════════════════════════════════════════════════════
# GET /api/metrics/charts/weekly-evolution
# ═══════════════════════════════════════════════════════════════════════

@router.get("/charts/weekly-evolution", response_model=WeeklyEvolutionResponse)
async def chart_weekly_evolution(
    data_inicio: str = Query(..., description="YYYY-MM-DD"),
    data_fim: str = Query(..., description="YYYY-MM-DD"),
    empresa_id: Optional[UUID] = Query(None),
    semanas: int = Query(DEFAULT_WEEKS, ge=1, le=12),
    current_user: Usuario = Depends(require_role(*ALL_ROLES)),
    db: AsyncSession = Depends(get_db),
):
    inicio = _parse_date(data_inicio)
    fim = _parse_date(data_fim)

    data = await get_weekly_evolution(db, inicio, fim, semanas, empresa_id, current_user)
    return WeeklyEvolutionResponse(semanas=data)


# ═══════════════════════════════════════════════════════════════════════
# GET /api/metrics/export/pdf
# ═══════════════════════════════════════════════════════════════════════

@router.get("/export/pdf")
@audit(action="METRICS_EXPORT_PDF", resource="METRICS", severity="INFO")
async def export_pdf(
    request: Request,
    data_inicio: str = Query(..., description="YYYY-MM-DD"),
    data_fim: str = Query(..., description="YYYY-MM-DD"),
    empresa_id: Optional[UUID] = Query(None),
    current_user: Usuario = Depends(require_role(*ALL_ROLES)),
    db: AsyncSession = Depends(get_db),
):
    inicio = _parse_date(data_inicio)
    fim = _parse_date(data_fim)

    metrics = await calculate_weekly_metrics(db, inicio, fim, empresa_id, current_user)

    try:
        from app.services.pdf_service import build_weekly_report

        report_data = {
            "report_id": f"metrics_{inicio.isoformat()}_{fim.isoformat()}",
            "company_name": metrics.get("empresa_nome") or "Todas as Empresas",
            "generated_by": current_user.full_name,
            "week_start": inicio.strftime("%d/%m/%Y"),
            "week_end": fim.strftime("%d/%m/%Y"),
            "active_people": metrics["pessoas_ativas"],
            "planned": metrics["ofc_programadas"],
            "realized": metrics["ofc_realizadas"],
            "positive": metrics["ofc_positivas"],
            "negative": metrics["ofc_negativas"],
            "adherence": metrics["aderencia_percentual"],
            "pct_seguro": metrics["percentual_seguro"],
            "pct_negativo": metrics["percentual_negativo"],
            "active_users": metrics["usuarios_ativos"],
            "avg_ofc_per_user": metrics["media_ofc_usuario"],
            "top_safe": metrics.get("top_comportamentos_seguros", []),
            "top_unsafe": metrics.get("top_comportamentos_negativos", []),
            "show_signature": True,
            "manager_name": current_user.full_name,
            "manager_role": current_user.role,
        }

        file_path = build_weekly_report(report_data)

        return StreamingResponse(
            file_path.open("rb"),
            media_type="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="{file_path.name}"',
                "X-Report-Id": report_data["report_id"],
            },
        )
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="WeasyPrint nao esta instalado. Execute: pip install weasyprint",
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erro ao gerar PDF: {str(e)}")


# ═══════════════════════════════════════════════════════════════════════
# GET /api/metrics/export/excel
# ═══════════════════════════════════════════════════════════════════════

@router.get("/export/excel")
@audit(action="METRICS_EXPORT_EXCEL", resource="METRICS", severity="INFO")
async def export_excel(
    request: Request,
    data_inicio: str = Query(..., description="YYYY-MM-DD"),
    data_fim: str = Query(..., description="YYYY-MM-DD"),
    empresa_id: Optional[UUID] = Query(None),
    current_user: Usuario = Depends(require_role(*ALL_ROLES)),
    db: AsyncSession = Depends(get_db),
):
    inicio = _parse_date(data_inicio)
    fim = _parse_date(data_fim)

    metrics = await calculate_weekly_metrics(db, inicio, fim, empresa_id, current_user)
    company_ranking = await get_company_ranking(db, inicio, fim, empresa_id, current_user)
    user_ranking = await get_user_ranking(db, inicio, fim, empresa_id, current_user)
    shift_dist = await get_shift_distribution(db, inicio, fim, empresa_id, current_user)
    behaviors = await get_top_behaviors(db, inicio, fim, empresa_id, current_user)

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="openpyxl nao esta instalado. Execute: pip install openpyxl",
        )

    wb = openpyxl.Workbook()

    # ── Aba 1: Indicadores ──
    ws = wb.active
    ws.title = "Indicadores"
    _write_indicators_sheet(ws, metrics, inicio, fim)

    # ── Aba 2: Ranking Empresas ──
    ws2 = wb.create_sheet("Ranking Empresas")
    _write_ranking_sheet(ws2, company_ranking, ["Empresa", "Total", "Positivas", "Negativas", "Aderencia"],
                         ["empresa_nome", "total", "positivas", "negativas", "aderencia"])

    # ── Aba 3: Ranking Usuarios ──
    ws3 = wb.create_sheet("Ranking Usuarios")
    _write_ranking_sheet(ws3, user_ranking,
                         ["Usuario", "Empresa", "Total", "Positivas", "Negativas"],
                         ["usuario_nome", "empresa_nome", "total", "positivas", "negativas"])

    # ── Aba 4: Por Turno ──
    ws4 = wb.create_sheet("Por Turno")
    _write_simple_sheet(ws4, shift_dist, ["Turno", "Quantidade"], ["name", "value"])

    # ── Aba 5: Comportamentos ──
    ws5 = wb.create_sheet("Comportamentos Seguros")
    _write_simple_sheet(ws5, behaviors["seguros"],
                        ["Comportamento", "Tipo", "Quantidade"],
                        ["comportamento", "tipo", "count"])

    ws6 = wb.create_sheet("Comportamentos Negativos")
    _write_simple_sheet(ws6, behaviors.get("negativos", behaviors.get("desvios", [])),
                        ["Comportamento", "Tipo", "Quantidade"],
                        ["comportamento", "tipo", "count"])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    file_name = f"metricas_{inicio.isoformat()}_{fim.isoformat()}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )


# ═══════════════════════════════════════════════════════════════════════
# GET /api/metrics/export/csv
# ═══════════════════════════════════════════════════════════════════════

@router.get("/export/csv")
@audit(action="METRICS_EXPORT_CSV", resource="METRICS", severity="INFO")
async def export_csv(
    request: Request,
    data_inicio: str = Query(..., description="YYYY-MM-DD"),
    data_fim: str = Query(..., description="YYYY-MM-DD"),
    empresa_id: Optional[UUID] = Query(None),
    current_user: Usuario = Depends(require_role(*ALL_ROLES)),
    db: AsyncSession = Depends(get_db),
):
    import csv

    inicio = _parse_date(data_inicio)
    fim = _parse_date(data_fim)

    metrics = await calculate_weekly_metrics(db, inicio, fim, empresa_id, current_user)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Indicador", "Valor", "Unidade", "Status"])
    for ind in metrics.get("indicadores", []):
        writer.writerow([ind["nome"], ind["valor"], ind["unidade"], ind.get("status", "")])

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={
            "Content-Disposition": f'attachment; filename=metricas_{inicio.isoformat()}_{fim.isoformat()}.csv',
        },
    )


# ═══════════════════════════════════════════════════════════════════════
# POST /api/metrics/cache/invalidate
# ═══════════════════════════════════════════════════════════════════════

@router.post("/cache/invalidate")
async def invalidate_cache(
    current_user: Usuario = Depends(require_role("gestor", "admin")),
):
    invalidate_metrics_cache()
    return {"message": "Cache de metricas invalidado com sucesso."}


# ═══════════════════════════════════════════════════════════════════════
# Helpers: geracao de planilhas Excel
# ═══════════════════════════════════════════════════════════════════════

def _make_excel_styles():
    """Factory para estilos Excel (importados sob demanda)."""
    try:
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        Font = None
        PatternFill = None
        Alignment = None
        Border = None
        Side = None
    
    return {
        "HEADER_FONT": Font(name="Calibri", size=11, bold=True, color="FFFFFF") if Font else None,
        "HEADER_FILL": PatternFill(start_color="1A1A1A", end_color="1A1A1A", fill_type="solid") if PatternFill else None,
        "HEADER_ALIGN": Alignment(horizontal="center", vertical="center") if Alignment else None,
        "CELL_ALIGN": Alignment(horizontal="left", vertical="center") if Alignment else None,
        "THIN_BORDER": Border(
            left=Side(style="thin", color="CCCCCC") if Side else None,
            right=Side(style="thin", color="CCCCCC") if Side else None,
            top=Side(style="thin", color="CCCCCC") if Side else None,
            bottom=Side(style="thin", color="CCCCCC") if Side else None,
        ) if Border else None,
    }


def _write_indicators_sheet(ws, metrics: dict, inicio: date, fim: date):
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = f"Relatorio de Metricas — {inicio.strftime('%d/%m/%Y')} a {fim.strftime('%d/%m/%Y')}"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1A1A1A")
    title_cell.alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:D2")
    ws["A2"].value = f"Gerado por: {metrics.get('gerado_por', '')} | Status: {metrics.get('status', '')}"
    ws["A2"].font = Font(name="Calibri", size=10, color="666666")
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["Indicador", "Valor", "Unidade", "Status"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER

    for i, ind in enumerate(metrics.get("indicadores", []), 5):
        ws.cell(row=i, column=1, value=ind["nome"]).border = _THIN_BORDER
        ws.cell(row=i, column=2, value=ind["valor"]).border = _THIN_BORDER
        ws.cell(row=i, column=3, value=ind["unidade"]).border = _THIN_BORDER
        ws.cell(row=i, column=4, value=ind.get("status", "")).border = _THIN_BORDER

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 12


def _write_ranking_sheet(ws, data: list, headers: list[str], keys: list[str]):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER

    for i, item in enumerate(data, 2):
        for col, key in enumerate(keys, 1):
            ws.cell(row=i, column=col, value=item.get(key, "")).border = _THIN_BORDER

    if get_column_letter:
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 22


def _write_simple_sheet(ws, data: list, headers: list[str], keys: list[str]):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER

    for i, item in enumerate(data, 2):
        for col, key in enumerate(keys, 1):
            ws.cell(row=i, column=col, value=item.get(key, "")).border = _THIN_BORDER

    if get_column_letter:
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 25
