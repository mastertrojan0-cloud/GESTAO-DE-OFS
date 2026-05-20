"""
excel_service.py — Exportacao Excel (openpyxl) do relatorio METRICAS DA SEMANA.
3 abas: Resumo, Consolidado, Registros Base.
Formatacao corporativa: cores, fontes, bordas.
"""
import io
import logging
from datetime import datetime
from typing import Optional
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference

from .pdf_service import gerar_nome_arquivo_excel

logger = logging.getLogger(__name__)

# ═══════════════════════════════════════════════════════════════════════
# Constantes de estilo corporativo
# ═══════════════════════════════════════════════════════════════════════

CORP_RED = "CC0000"
CORP_BLACK = "1A1A1A"
CORP_GRAY = "F5F5F5"
CORP_WHITE = "FFFFFF"
CORP_GREEN = "22C55E"
CORP_YELLOW = "EAB308"
CORP_RED_DANGER = "EF4444"
CORP_DARK_GRAY = "666666"
CORP_BLUE = "3B82F6"

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

HEADER_FILL = PatternFill(start_color=CORP_RED, end_color=CORP_RED, fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color=CORP_WHITE)
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color=CORP_BLACK)
SUBTITLE_FONT = Font(name="Calibri", size=11, bold=True, color=CORP_BLACK)
BODY_FONT = Font(name="Calibri", size=10, color=CORP_BLACK)
BODY_FONT_BOLD = Font(name="Calibri", size=10, bold=True, color=CORP_BLACK)
META_FONT = Font(name="Calibri", size=9, color=CORP_DARK_GRAY)
TOTAL_FILL = PatternFill(start_color="E5E5E5", end_color="E5E5E5", fill_type="solid")
ALT_FILL = PatternFill(start_color=CORP_GRAY, end_color=CORP_GRAY, fill_type="solid")

STATUS_FILLS = {
    "OK": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
    "ATENCAO": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
    "ALERTA": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
}

STATUS_FONTS = {
    "OK": Font(name="Calibri", size=10, bold=True, color=CORP_GREEN),
    "ATENCAO": Font(name="Calibri", size=10, bold=True, color="A16207"),
    "ALERTA": Font(name="Calibri", size=10, bold=True, color=CORP_RED_DANGER),
}

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

EXPORTS_DIR = Path(os.environ.get("EXPORTS_DIR", "/app/reports"))


# ═══════════════════════════════════════════════════════════════════════
# Helpers de formatacao
# ═══════════════════════════════════════════════════════════════════════

def _style_header_row(ws, row: int, col_count: int):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def _style_data_cell(ws, row: int, col: int, alignment: Alignment = CENTER):
    cell = ws.cell(row=row, column=col)
    cell.font = BODY_FONT
    cell.alignment = alignment
    cell.border = THIN_BORDER


def _apply_alt_rows(ws, start_row: int, end_row: int, col_count: int):
    for r in range(start_row, end_row + 1):
        if (r - start_row) % 2 == 1:
            for c in range(1, col_count + 1):
                ws.cell(row=r, column=c).fill = ALT_FILL


def _auto_width(ws, col_count: int, min_width: int = 10, max_width: int = 45):
    for col in range(1, col_count + 1):
        max_len = min_width
        for row in ws.iter_rows(min_col=col, max_col=col, values_only=False):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)) + 2, max_width))
        ws.column_dimensions[get_column_letter(col)].width = max_len


def _status(value: float, indicator: str) -> str:
    from .pdf_service import evaluate_status
    if value is None:
        return ""
    return evaluate_status(indicator, value)


# ═══════════════════════════════════════════════════════════════════════
# Aba 1: RESUMO
# ═══════════════════════════════════════════════════════════════════════

def _build_resumo_sheet(wb: Workbook, data: dict) -> None:
    ws = wb.active
    ws.title = "Resumo"
    ws.sheet_properties.tabColor = CORP_RED

    ws.merge_cells("A1:F1")
    title_cell = ws.cell(row=1, column=1, value="METRICAS DA SEMANA")
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A2:F2")
    week_start = data.get("week_start", "")
    week_end = data.get("week_end", "")
    ws.cell(row=2, column=1,
            value=f"Semana {data.get('semana', '-')}/{data.get('ano', '-')}  |  {week_start} a {week_end}").font = META_FONT

    ws.merge_cells("A3:F3")
    ws.cell(row=3, column=1,
            value=f"Empresa: {data.get('company_name', 'Security Dynamics Ltda.')}").font = META_FONT

    ws.merge_cells("A4:F4")
    gerado = data.get("generated_by", "")
    ws.cell(row=4, column=1,
            value=f"Gerado por: {gerado}  |  {datetime.now().strftime('%d/%m/%Y %H:%M')}").font = META_FONT

    r = 6
    ws.merge_cells(f"A{r}:F{r}")
    ws.cell(row=r, column=1, value="INDICADORES PRINCIPAIS").font = SUBTITLE_FONT
    ws.cell(row=r, column=1).fill = PatternFill(start_color=CORP_GRAY, end_color=CORP_GRAY, fill_type="solid")

    r += 1
    headers = ["Indicador", "Valor", "Status", "Meta", "", ""]
    for c, h in enumerate(headers, 1):
        ws.cell(row=r, column=c, value=h)
    _style_header_row(ws, r, 6)
    ws.merge_cells(f"E{r}:F{r}")
    ws.cell(row=r, column=5, value="Observacoes").font = HEADER_FONT
    ws.cell(row=r, column=5).alignment = CENTER

    indicators = [
        ("Pessoas Ativas", data.get("active_people", 0), "", ">= 30", "Quantidade de pessoas no contrato"),
        ("OFS Programadas", data.get("planned", 0), "", "Meta da semana", "pessoas_ativas x meta_por_pessoa"),
        ("OFS Realizadas", data.get("realized", 0), "", "-", "Registros nao cancelados"),
        ("Aderencia %", f"{data.get('adherence', 0):.1f}%",
         _status(data.get("adherence", 0), "aderencia"), ">= 100%",
         "(realizadas / programadas) x 100"),
        ("Positivas", data.get("positive", 0), "", "-", "Registros tipo Positivo/Seguro"),
        ("Negativas", data.get("negative", 0), "", "-", "Registros tipo Negativo/Inseguro"),
        ("% Seguro", f"{data.get('pct_seguro', 0):.1f}%",
         _status(data.get("pct_seguro", 0), "pct_seguro"), ">= 80%",
         "(positivas / realizadas) x 100"),
        ("% Desvio", f"{data.get('pct_desvio', 0):.1f}%",
         _status(data.get("pct_desvio", 0), "pct_desvio"), "<= 15%",
         "(negativas / realizadas) x 100"),
        ("Usuarios Ativos", data.get("active_users", 0), "", "-", "Usuarios com ao menos 1 registro"),
        ("Media OFS / Usuario", f"{data.get('avg_ofc_per_user', 0):.1f}", "", ">= 15",
         "realizadas / usuarios_ativos"),
    ]

    r += 1
    for i, (ind, val, status_val, meta, obs) in enumerate(indicators):
        row = r + i
        ws.cell(row=row, column=1, value=ind).font = BODY_FONT_BOLD
        ws.cell(row=row, column=1).alignment = LEFT
        ws.cell(row=row, column=2, value=val).font = BODY_FONT
        ws.cell(row=row, column=2).alignment = CENTER
        ws.cell(row=row, column=3, value=status_val).font = STATUS_FONTS.get(status_val, BODY_FONT)
        ws.cell(row=row, column=3).alignment = CENTER
        if status_val in STATUS_FILLS:
            ws.cell(row=row, column=3).fill = STATUS_FILLS[status_val]
        ws.cell(row=row, column=4, value=meta).font = META_FONT
        ws.cell(row=row, column=4).alignment = CENTER
        ws.merge_cells(f"E{row}:F{row}")
        ws.cell(row=row, column=5, value=obs).font = META_FONT
        ws.cell(row=row, column=5).alignment = LEFT
        for c in range(1, 7):
            ws.cell(row=row, column=c).border = THIN_BORDER
            if i % 2 == 1 and c not in (3,):
                ws.cell(row=row, column=c).fill = ALT_FILL

    _auto_width(ws, 6, min_width=14)
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["F"].width = 15


# ═══════════════════════════════════════════════════════════════════════
# Aba 2: CONSOLIDADO
# ═══════════════════════════════════════════════════════════════════════

def _build_consolidado_sheet(wb: Workbook, data: dict) -> None:
    ws = wb.create_sheet("Consolidado")
    ws.sheet_properties.tabColor = "3B82F6"

    ws.merge_cells("A1:H1")
    ws.cell(row=1, column=1, value="TABELA CONSOLIDADA POR EMPRESA").font = SUBTITLE_FONT
    ws.cell(row=1, column=1).fill = PatternFill(start_color=CORP_GRAY, end_color=CORP_GRAY, fill_type="solid")

    headers = ["Empresa", "Pessoas Ativas", "Programadas", "Realizadas",
               "Positivas", "Negativas", "Aderencia %", "Status"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=2, column=c, value=h)
    _style_header_row(ws, 2, 8)

    consolidated = data.get("consolidated", [])
    if not consolidated:
        ws.cell(row=3, column=1, value="Nenhum dado disponivel").font = META_FONT
        return

    for i, row_data in enumerate(consolidated):
        r = 3 + i
        vals = [
            row_data.get("company", "-"),
            row_data.get("active_people", 0),
            row_data.get("planned", 0),
            row_data.get("realized", 0),
            row_data.get("positive", 0),
            row_data.get("negative", 0),
            f"{row_data.get('adherence', 0):.1f}%",
            row_data.get("status", "-"),
        ]
        for c, val in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=val)
            _style_data_cell(ws, r, c, CENTER if c > 1 else LEFT)
            if c == 5:
                ws.cell(row=r, column=c).font = Font(name="Calibri", size=10, color=CORP_GREEN)
            elif c == 6:
                ws.cell(row=r, column=c).font = Font(name="Calibri", size=10, color=CORP_RED_DANGER)
        status_val = row_data.get("status", "")
        if status_val in STATUS_FILLS:
            ws.cell(row=r, column=8).fill = STATUS_FILLS[status_val]
        if status_val in STATUS_FONTS:
            ws.cell(row=r, column=8).font = STATUS_FONTS[status_val]

    _apply_alt_rows(ws, 3, 2 + len(consolidated), 8)

    total_r = 3 + len(consolidated)
    total_data = [
        "TOTAL",
        sum(r.get("active_people", 0) for r in consolidated),
        sum(r.get("planned", 0) for r in consolidated),
        sum(r.get("realized", 0) for r in consolidated),
        sum(r.get("positive", 0) for r in consolidated),
        sum(r.get("negative", 0) for r in consolidated),
        f"{(sum(r.get('realized', 0) for r in consolidated) / max(sum(r.get('planned', 0) for r in consolidated), 1) * 100):.1f}%",
        "",
    ]
    for c, val in enumerate(total_data, 1):
        ws.cell(row=total_r, column=c, value=val)
        ws.cell(row=total_r, column=c).font = BODY_FONT_BOLD
        ws.cell(row=total_r, column=c).fill = TOTAL_FILL
        ws.cell(row=total_r, column=c).alignment = CENTER if c > 1 else LEFT
        ws.cell(row=total_r, column=c).border = Border(
            top=Side(style="medium", color=CORP_BLACK),
            bottom=Side(style="medium", color=CORP_BLACK),
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
        )

    _auto_width(ws, 8, min_width=12)
    ws.column_dimensions["A"].width = 22


# ═══════════════════════════════════════════════════════════════════════
# Aba 3: REGISTROS BASE
# ═══════════════════════════════════════════════════════════════════════

def _build_registros_sheet(wb: Workbook, data: dict) -> None:
    ws = wb.create_sheet("Registros Base")
    ws.sheet_properties.tabColor = "999999"

    ws.merge_cells("A1:L1")
    ws.cell(row=1, column=1, value="REGISTROS BASE - DADOS BRUTOS DA SEMANA").font = SUBTITLE_FONT
    ws.cell(row=1, column=1).fill = PatternFill(start_color=CORP_GRAY, end_color=CORP_GRAY, fill_type="solid")

    headers = [
        "No OFS", "Data", "Hora", "Gerado por", "Empresa Observada",
        "Observado", "Atividade", "Local", "Turno", "Tipo",
        "Comportamento", "Status",
    ]
    for c, h in enumerate(headers, 1):
        ws.cell(row=2, column=c, value=h)
    _style_header_row(ws, 2, 12)

    records = data.get("base_records", [])
    if not records:
        ws.cell(row=3, column=1, value="Nenhum registro no periodo").font = META_FONT
        return

    for i, rec in enumerate(records):
        r = 3 + i
        vals = [
            rec.get("sequential_number", "-"),
            rec.get("record_date", "-"),
            rec.get("record_time", "-"),
            rec.get("generated_by_name", "-"),
            rec.get("company_name", "-"),
            rec.get("observed_name", "-"),
            rec.get("activity_observed", "-"),
            rec.get("location_observed", "-"),
            rec.get("shift", "-"),
            rec.get("type", "-"),
            rec.get("behavior_observed", "-"),
            rec.get("status", "ativo"),
        ]
        for c, val in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=val)
            _style_data_cell(ws, r, c, CENTER if c <= 2 or c == 10 else LEFT)

        tipo = rec.get("type", "")
        if "OFS" in tipo.upper() or "POSITIVO" in tipo.upper() or "SEGURO" in tipo.upper():
            ws.cell(row=r, column=10).fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
            ws.cell(row=r, column=10).font = Font(name="Calibri", size=10, color=CORP_GREEN, bold=True)
        elif "OFS" in tipo.upper() or "NEGATIVO" in tipo.upper() or "INSEGURO" in tipo.upper():
            ws.cell(row=r, column=10).fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            ws.cell(row=r, column=10).font = Font(name="Calibri", size=10, color=CORP_RED_DANGER, bold=True)

    _apply_alt_rows(ws, 3, 2 + len(records), 12)
    _auto_width(ws, 12, min_width=8)
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["G"].width = 25
    ws.column_dimensions["H"].width = 18
    ws.column_dimensions["K"].width = 45

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:L{2 + len(records)}"


# ═══════════════════════════════════════════════════════════════════════
# Funcoes principais de exportacao
# ═══════════════════════════════════════════════════════════════════════

def export_weekly_excel(data: dict) -> Path:
    """Gera arquivo Excel do relatorio METRICAS DA SEMANA. Retorna o caminho do .xlsx."""
    wb = Workbook()
    _build_resumo_sheet(wb, data)
    _build_consolidado_sheet(wb, data)
    _build_registros_sheet(wb, data)

    file_name = gerar_nome_arquivo_excel("weekly_metrics", {
        "ano": data.get("ano", datetime.now().year),
        "semana": data.get("semana", 0),
    })

    output_dir = EXPORTS_DIR / "excel" / str(data.get("ano", datetime.now().year))
    output_dir.mkdir(parents=True, exist_ok=True)
    file_path = output_dir / file_name

    wb.save(str(file_path))
    logger.info(f"Excel gerado: {file_path} ({file_path.stat().st_size} bytes)")
    return file_path


def export_weekly_excel_bytes(data: dict) -> bytes:
    """Retorna o arquivo Excel como bytes (para download direto via API)."""
    wb = Workbook()
    _build_resumo_sheet(wb, data)
    _build_consolidado_sheet(wb, data)
    _build_registros_sheet(wb, data)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
